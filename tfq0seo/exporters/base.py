"""Export analysis results through a shared, presentation-only adapter."""

import csv
import json
import math
import re
from string import Formatter
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional
from urllib.parse import urlsplit

from jinja2 import Environment, FileSystemLoader, TemplateNotFound, select_autoescape
from ..core.report_contracts import validate_report

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class ExportManager:
    """Manages JSON, CSV, Excel and HTML exports without changing source results."""

    ANALYZERS = ('seo', 'content', 'technical', 'performance', 'links')

    def __init__(self, config: Optional[Any] = None):
        self.config = dict(vars(config) if hasattr(config, '__dict__') else config or {})
        self.output_dir = Path(self.config.get('output_directory', './reports'))
        self.output_dir.mkdir(parents=True, exist_ok=True)
        template_dir = Path(__file__).parent.parent / 'templates'
        self.jinja_env = Environment(
            loader=FileSystemLoader(str(template_dir)),
            autoescape=select_autoescape(['html', 'xml'], default_for_string=True),
        )
        self.jinja_env.filters.update(number=self._display_number, score_color=self._score_color,
                                     safe_url=self._safe_url, json_text=self._json_text)
        self.jinja_env.policies['json.dumps_kwargs'] = {'default': str, 'sort_keys': False}

    @staticmethod
    def _number(value):
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            try:
                return value if math.isfinite(value) else None
            except OverflowError:
                pass
        return None

    @classmethod
    def _display_number(cls, value, digits=0, suffix=''):
        number = cls._number(value)
        if number is None:
            return 'Not measured'
        return (str(round(number)) if digits == 0 else f'{number:.{digits}f}') + suffix

    @classmethod
    def _score_color(cls, value):
        number = cls._number(value)
        if number is None:
            return 'var(--color-gray, #666)'
        return 'var(--color-pass)' if number >= 80 else 'var(--color-average)' if number >= 50 else 'var(--color-fail)'

    @staticmethod
    def _safe_url(value):
        """Only allow complete HTTP(S) destinations in clickable report links."""
        if not isinstance(value, str) or any(ord(char) < 32 for char in value):
            return ''
        try:
            parsed = urlsplit(value)
            if parsed.scheme.lower() in ('http', 'https') and parsed.hostname:
                return value
        except ValueError:
            pass
        return ''

    @staticmethod
    def _json_text(value):
        """Return plain JSON text for autoescaped HTML evidence, never trusted markup."""
        return json.dumps(value, indent=2, ensure_ascii=False, default=str)

    @staticmethod
    def _cell(value):
        """Keep cells scalar and prevent untrusted strings becoming formulas."""
        if value is None or isinstance(value, (bool, int, float)):
            return value
        if not isinstance(value, str):
            value = json.dumps(value, ensure_ascii=False, default=str)
        value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', value)
        if value.lstrip().startswith(('=', '+', '-', '@')) or value.startswith(('\t', '\r', '\n')):
            value = "'" + value
        return value

    @staticmethod
    def _records(value):
        return [item for item in value if isinstance(item, dict)] if isinstance(value, list) else []

    @classmethod
    def _page_records(cls, data):
        pages = data.get('pages')
        if isinstance(pages, dict):
            records = cls._records(pages.get('detailed')) or cls._records(pages.get('summary'))
            return records + cls._records(pages.get('failed')) + cls._records(pages.get('skipped'))
        if isinstance(pages, list):
            return cls._records(pages)
        return [data]

    @staticmethod
    def _recommendation_text(rec):
        if isinstance(rec, dict):
            return str(rec.get('recommendation') or rec.get('description') or rec.get('title') or '')
        return str(rec)

    @classmethod
    def _recommendations(cls, data):
        recs = data.get('recommendations', [])
        if isinstance(recs, dict):
            recs = recs.get('specific', [])
        return recs if isinstance(recs, list) else []

    @classmethod
    def _scores(cls, data):
        scores = data.get('scores', {})
        overall = cls._number(scores.get('overall', data.get('overall_score')))
        categories = scores.get('categories', data.get('category_scores'))
        if not isinstance(categories, dict):
            categories = {name: data.get(name, {}).get('score') for name in cls.ANALYZERS}
        return overall, {name: cls._number(value) for name, value in categories.items()}

    @classmethod
    def _issues(cls, data):
        issues = data.get('issues', [])
        if isinstance(issues, dict):
            issues = issues.get('aggregated', issues.get('top_issues', []))
        issues = cls._records(issues)
        return sorted(issues, key=lambda issue: {'critical': 0, 'warning': 1, 'notice': 2}.get(issue.get('severity'), 3))

    @classmethod
    def _issue_counts(cls, data, issues=None):
        issues = cls._issues(data) if issues is None else issues
        counts = {severity: sum(issue.get('severity') == severity for issue in issues)
                  for severity in ('critical', 'warning', 'notice')}
        counts.update(total=len(issues), unique=len(issues))
        issue_data = data.get('issues')
        supplied = issue_data.get('counts', {}) if isinstance(issue_data, dict) else data.get('issue_counts', {})
        counts.update({key: value for key, value in supplied.items() if cls._number(value) is not None})
        return counts

    def export(self, data: Dict[str, Any], format: str, output_path: Optional[str] = None) -> str:
        if format not in ('json', 'csv', 'xlsx', 'html'):
            raise ValueError(f'Unsupported export format: {format}')
        validate_report(data)
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S_%f')
        if output_path:
            output_file = Path(output_path)
        else:
            pattern = self.config.get('filename_pattern', '{domain}_{timestamp}_{format}')
            for _, field, spec, conversion in Formatter().parse(pattern):
                if field is not None and (field not in ('domain', 'timestamp', 'format') or spec or conversion):
                    raise ValueError('filename_pattern supports only {domain}, {timestamp}, and {format}')
            records = self._page_records(data)
            try:
                domain = urlsplit(records[0].get('url', '')).hostname if records else None
            except ValueError:
                domain = None
            domain = re.sub(r'[^A-Za-z0-9.-]', '_', domain or 'site')[:120]
            filename = pattern.format(domain=domain, timestamp=timestamp, format=format)
            device_name = filename.split('.')[0].rstrip(' ').upper()
            reserved = device_name in {'CON', 'PRN', 'AUX', 'NUL', 'CONIN$', 'CONOUT$'} or re.fullmatch(
                r'(?:COM|LPT)[1-9¹²³]', device_name)
            if (not filename or filename in ('.', '..') or Path(filename).name != filename
                    or re.search(r'[<>:"/\\|?*\x00-\x1f]', filename) or reserved):
                raise ValueError('filename_pattern must produce a portable filename without directories or reserved names')
            if not filename.endswith('.' + format):
                filename += '.' + format
            output_file = self.output_dir / filename
        output_file.parent.mkdir(parents=True, exist_ok=True)
        return getattr(self, f'export_{format}')(data, output_file)

    def export_json(self, data: Dict[str, Any], output_file: Path) -> str:
        validate_report(data)
        serialized = json.dumps(data, indent=2, default=str, ensure_ascii=False, allow_nan=False)
        with open(output_file, 'w', encoding='utf-8') as stream:
            stream.write(serialized)
        return str(output_file)

    def export_csv(self, data: Dict[str, Any], output_file: Path) -> str:
        validate_report(data)
        rows = [self._flatten_page_data(page) for page in self._page_records(data)]
        rows = rows or [{'error': 'No data to export'}]
        fieldnames = sorted({key for row in rows for key in row})
        with open(output_file, 'w', newline='', encoding='utf-8') as stream:
            writer = csv.DictWriter(stream, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows({key: self._cell(value) for key, value in row.items()} for row in rows)
        return str(output_file)

    def _flatten_page_data(self, page_data: Dict[str, Any]) -> Dict[str, Any]:
        flat = {key: page_data.get(key) for key in ('url', 'status_code', 'load_time', 'error', 'status')}
        flat['overall_score'] = self._number(page_data.get('overall_score', page_data.get('score')))
        for name in self.ANALYZERS:
            analyzer = page_data.get(name, {})
            flat[f'{name}_score'] = self._number(analyzer.get('score'))
            flat[f'{name}_error'] = analyzer.get('error')
        counts = page_data.get('issue_counts', page_data.get('issues', {}))
        if isinstance(counts, dict):
            flat.update({f'issues_{key}': value for key, value in counts.items()})
        seo_data = page_data.get('seo', {}).get('data', {})
        for key in ('title', 'description'):
            value = seo_data.get(key, '')
            flat[key] = value.get('text', '') if isinstance(value, dict) else value
        content = page_data.get('content', {}).get('data', {})
        flat['word_count'] = content.get('metrics', {}).get('word_count', content.get('word_count'))
        flat['flesch_reading_ease'] = content.get('readability', {}).get('flesch_reading_ease', content.get('flesch_reading_ease'))
        performance = page_data.get('performance', {}).get('data', {})
        for key in ('total_resources', 'content_size_mb'):
            flat[key] = performance.get(key, performance.get('metrics', {}).get(key))
        return flat

    def export_xlsx(self, data: Dict[str, Any], output_file: Path) -> str:
        validate_report(data)
        if not OPENPYXL_AVAILABLE:
            return self.export_csv(data, output_file.with_suffix('.csv'))
        workbook = Workbook()
        summary = workbook.active
        summary.title = 'Summary'
        self._write_summary_sheet(summary, data)
        self._write_issues_sheet(workbook.create_sheet('Issues'), data)
        self._write_pages_sheet(workbook.create_sheet('Pages'), self._page_records(data))
        self._write_recommendations_sheet(workbook.create_sheet('Recommendations'), self._recommendations(data))
        workbook.save(output_file)
        return str(output_file)

    def _append_row(self, sheet, values):
        sheet.append([self._cell(value) for value in values])

    @staticmethod
    def _style_header(sheet):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')
        sheet.freeze_panes = 'A2'

    def _write_summary_sheet(self, sheet, data):
        overall, categories = self._scores(data)
        self._append_row(sheet, ['SEO Analysis Report'])
        self._append_row(sheet, [])
        self._append_row(sheet, ['Overall Score', overall if overall is not None else 'Not measured'])
        self._append_row(sheet, ['Assessment', data.get('status', '')])
        for name, score in categories.items():
            self._append_row(sheet, [name.title(), score if score is not None else 'Not measured'])
        counts = self._issue_counts(data)
        for severity in ('critical', 'warning', 'notice', 'total'):
            self._append_row(sheet, [severity.title() + ' Issues', counts[severity]])
        if data.get('error'):
            self._append_row(sheet, ['Error', data['error']])
        self._style_header(sheet)
        sheet.column_dimensions['A'].width = 24
        sheet.column_dimensions['B'].width = 30

    def _write_issues_sheet(self, sheet, data):
        self._append_row(sheet, ['Severity', 'Category', 'Message', 'URL', 'Occurrences', 'Affected Pages'])
        issues = self._issues(data)
        if not issues and 'pages' in data:
            issues = [dict(issue, url=page.get('url')) for page in self._page_records(data) for issue in self._issues(page)]
        for issue in issues:
            self._append_row(sheet, [issue.get('severity'), issue.get('category'), issue.get('message'),
                                     issue.get('url', data.get('url')), issue.get('count', 1),
                                     issue.get('pages', issue.get('pages_affected'))])
        self._style_header(sheet)

    def _write_pages_sheet(self, sheet, pages: List[Dict[str, Any]]):
        rows = [self._flatten_page_data(page) for page in pages]
        keys = ['url', 'overall_score', 'status_code', 'load_time']
        keys += sorted({key for row in rows for key in row} - set(keys))
        self._append_row(sheet, [key.replace('_', ' ').title() for key in keys])
        for row in rows:
            self._append_row(sheet, [row.get(key) for key in keys])
        self._style_header(sheet)

    def _write_recommendations_sheet(self, sheet, recommendations):
        if isinstance(recommendations, dict):
            recommendations = recommendations.get('specific', [])
        self._append_row(sheet, ['Recommendation', 'Priority', 'Category', 'Impact', 'Effort'])
        for rec in recommendations:
            metadata = rec if isinstance(rec, dict) else {}
            self._append_row(sheet, [self._recommendation_text(rec)] + [metadata.get(key) for key in ('priority', 'category', 'impact', 'effort')])
        self._style_header(sheet)

    def export_html(self, data: Dict[str, Any], output_file: Path) -> str:
        validate_report(data)
        requested = self.config.get('html_template')
        template_names = {'report': 'report.html', 'enhanced': 'enhanced_report.html', 'optimized': 'optimized_report.html'}
        if requested is not None and requested not in template_names:
            raise ValueError(f'Unsupported HTML template: {requested}')
        default = 'optimized' if 'pages' in data or 'aggregated_issues' in data else 'report'
        try:
            template = self.jinja_env.get_template(template_names[requested or default])
        except TemplateNotFound:
            try:
                template = self.jinja_env.get_template('report.html')
            except TemplateNotFound:
                template = self.jinja_env.from_string(self._get_basic_html_template())
        html = template.render(**self._prepare_html_data(data))
        with open(output_file, 'w', encoding='utf-8') as stream:
            stream.write(html)
        return str(output_file)

    def _prepare_html_data(self, data: Dict[str, Any]) -> Dict[str, Any]:
        overall, categories = self._scores(data)
        issues = self._issues(data)
        issue_data = data.get('issues', {}) if isinstance(data.get('issues'), dict) else {}
        counts = self._issue_counts(data, issues)
        recs = self._recommendations(data)
        recommendations = data.get('recommendations', {}) if isinstance(data.get('recommendations'), dict) else {}
        pages = self._page_records(data)
        summaries = []
        for page in pages:
            count_data = page.get('issue_counts', page.get('issues', {}))
            issue_count = count_data.get('total', 0) if isinstance(count_data, dict) else len(count_data)
            summaries.append({'url': str(page.get('url', '')), 'score': self._number(page.get('overall_score', page.get('score'))),
                              'status_code': page.get('status_code'), 'load_time': self._number(page.get('load_time')),
                              'issue_count': issue_count, 'error': page.get('error'), 'status': page.get('status')})
        summary = data.get('summary', {})
        executive_defaults = {
            'overview': {'total_pages_analyzed': summary.get('total_pages', len(pages)),
                         'successful_pages': summary.get('successful_pages', sum(not page.get('error') for page in pages)),
                         'overall_health': 'Not measured' if overall is None else 'Good' if overall >= 80 else 'Needs improvement'},
            'key_metrics': {'critical_issues': counts['critical']}, 'quick_wins': [],
        }
        supplied_executive = recommendations.get('executive', data.get('executive_summary')) or {}
        executive = {**executive_defaults, **supplied_executive}
        for key in ('overview', 'key_metrics'):
            executive[key] = {**executive_defaults[key], **supplied_executive.get(key, {})}
        performance = dict(data.get('performance', {}).get('data', {}))
        performance.update(performance.get('metrics', {}))
        performance.setdefault('html_fetch_time', data.get('load_time'))
        content = dict(data.get('content', {}).get('data', {}))
        content.update(content.get('metrics', {}))
        content.update(content.get('readability', {}))
        content.setdefault('average_grade_level', content.get('consensus_grade'))
        content.setdefault('image_count', content.get('image_analysis', {}).get('total_images'))
        technical = dict(data.get('technical', {}).get('data', {}))
        technical.setdefault('has_viewport', technical.get('mobile', {}).get('viewport_configured'))
        technical.setdefault('compression', technical.get('performance', {}).get('compression'))
        seo = dict(data.get('seo', {}).get('data', {}))
        for key in ('title', 'description'):
            if isinstance(seo.get(key), dict):
                seo[key] = seo[key].get('text', '')
        scoring = data.get('scoring') if isinstance(data.get('scoring'), dict) else {}
        scoring_policies = [scoring] if scoring else []
        if not scoring_policies:
            for page in pages:
                policy = page.get('scoring')
                if isinstance(policy, dict) and policy and policy not in scoring_policies:
                    scoring_policies.append(policy)
        versions = data.get('scores', {}).get('scoring_versions', [])
        scoring_versions = list(versions) if isinstance(versions, list) else []
        for policy in scoring_policies:
            if policy.get('version') is not None and policy['version'] not in scoring_versions:
                scoring_versions.append(policy['version'])
        rule_coverage = data.get('rule_coverage')
        if not isinstance(rule_coverage, dict):
            coverage = data.get('coverage', {})
            rule_coverage = coverage.get('rules', {}) if isinstance(coverage, dict) else {}
        if not isinstance(rule_coverage, dict):
            rule_coverage = {}
        enhanced_recommendations = recommendations.get('specific', data.get('enhanced_recommendations', []))
        enhanced_recommendations = [rec if isinstance(rec, dict) else {'title': rec, 'description': rec}
                                    for rec in enhanced_recommendations]
        aggregation_stats = issue_data.get('stats', data.get('aggregation_stats', {}))
        if self._number(aggregation_stats.get('reduction_ratio')) is None:
            aggregation_stats = {}
        return {
            'data': data,
            'url': data.get('url', 'Multiple Pages'), 'timestamp': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'error': data.get('error'), 'overall_score': overall, 'category_scores': categories,
            'report_status': data.get('status'),
            'scoring_policies': scoring_policies, 'scoring_versions': scoring_versions,
            'rule_coverage': rule_coverage,
            'score_source': data.get('scores', {}).get('source', data.get('coverage', {}).get('measurement_source')),
            'analysis_errors': [{'analyzer': name, 'error': data[name]['error']} for name in self.ANALYZERS
                                if isinstance(data.get(name), dict) and data[name].get('error')],
            'score_chart_values': [overall, 100 - overall] if overall is not None else [],
            'issues': issues, 'issue_counts': counts,
            'issues_by_severity': {severity: [issue for issue in issues if issue.get('severity') == severity]
                                   for severity in ('critical', 'warning', 'notice')},
            'recommendations': [self._recommendation_text(rec) for rec in recs],
            'enhanced_recommendations': enhanced_recommendations,
            'executive_summary': executive,
            'aggregated_issues': issue_data.get('aggregated', data.get('aggregated_issues', [])),
            'aggregation_stats': aggregation_stats,
            'performance_metrics': data.get('performance_metrics', {}),
            'pages': pages, 'pages_summary': summaries, 'pages_truncated': data.get('pages_truncated', False),
            'pages_truncated_count': data.get('pages_truncated_count', 0), 'summary': summary,
            'seo_data': seo, 'content_data': content, 'technical_data': technical,
            'performance_data': performance,
        }

    def _get_basic_html_template(self) -> str:
        return '''<!DOCTYPE html><html lang="en"><head><meta charset="UTF-8">
<title>SEO Analysis Report</title></head><body><h1>SEO Analysis Report</h1>
<p>{{ url }}</p><p>Overall score: {{ overall_score|number }}</p>
{% if error %}<p>Analysis failed: {{ error }}</p>{% endif %}
{% for issue in issues %}<p>{{ issue.category }}: {{ issue.message }}</p>{% endfor %}
{% for rec in recommendations %}<p>{{ rec }}</p>{% endfor %}
<table>{% for page in pages_summary %}<tr><td>{{ page.url }}</td><td>{{ page.score|number }}</td>
<td>{{ page.error or '' }}</td></tr>{% endfor %}</table></body></html>'''
