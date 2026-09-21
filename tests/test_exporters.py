"""Exercise exports using application results and untrusted audited content."""

import asyncio
import copy
import csv
import json
import re
from pathlib import Path

import pytest
from bs4 import BeautifulSoup
from jinja2 import FileSystemLoader
from openpyxl import load_workbook

from tfq0seo.core.app import SEOAnalyzer
from tfq0seo.exporters.base import ExportManager


@pytest.mark.parametrize('format', ['json', 'csv', 'html', 'xlsx'])
@pytest.mark.parametrize('existing', [False, True])
def test_failed_atomic_replace_preserves_destination_and_removes_temporary_files(
        tmp_path, monkeypatch, format, existing):
    exporter = ExportManager({'output_directory': str(tmp_path)})
    destination = tmp_path / ('report.' + format)
    if existing:
        destination.write_bytes(b'previous report')

    def fail_replace(source, target):
        assert Path(source).parent == destination.parent
        assert Path(source).stat().st_size > 0
        assert Path(target) == destination
        raise OSError('replacement denied')

    monkeypatch.setattr('tfq0seo.exporters.base.os.replace', fail_replace)
    with pytest.raises(OSError, match='replacement denied'):
        exporter.export({'url': 'https://example.test'}, format, str(destination))
    assert destination.read_bytes() == b'previous report' if existing else not destination.exists()
    assert list(tmp_path.iterdir()) == ([destination] if existing else [])


@pytest.mark.parametrize('format', ['json', 'csv', 'html', 'xlsx'])
def test_partial_writer_failure_preserves_previous_report(tmp_path, monkeypatch, format):
    exporter = ExportManager({'output_directory': str(tmp_path)})
    destination = tmp_path / ('report.' + format)
    destination.write_bytes(b'previous report')
    if format == 'json':
        def fail_json(data, stream, **kwargs):
            stream.write('{"partial":')
            raise OSError('writer failed')
        monkeypatch.setattr('tfq0seo.exporters.base.json.dump', fail_json)
    elif format == 'csv':
        def fail_rows(writer, rows):
            writer.writerow(next(iter(rows)))
            raise OSError('writer failed')
        monkeypatch.setattr(csv.DictWriter, 'writerows', fail_rows)
    elif format == 'html':
        def fail_html(*args, **kwargs):
            yield '<html>partial'
            raise OSError('writer failed')
        template = exporter.jinja_env.get_template('report.html')
        monkeypatch.setattr(template, 'generate', fail_html)
    else:
        def fail_workbook(workbook, stream):
            stream.write(b'partial zip')
            raise OSError('writer failed')
        monkeypatch.setattr('tfq0seo.exporters.base.Workbook.save', fail_workbook)
    with pytest.raises(OSError, match='writer failed'):
        getattr(exporter, 'export_' + format)({'url': 'https://example.test'}, destination)
    assert destination.read_bytes() == b'previous report'
    assert list(tmp_path.iterdir()) == [destination]


def test_atomic_csv_fallback_preserves_existing_xlsx(tmp_path, monkeypatch):
    exporter = ExportManager({'output_directory': str(tmp_path)})
    destination = tmp_path / 'report.xlsx'
    destination.write_bytes(b'previous workbook')
    monkeypatch.setattr('tfq0seo.exporters.base.OPENPYXL_AVAILABLE', False)
    result = exporter.export({'url': 'https://example.test'}, 'xlsx', str(destination))
    assert result == str(destination.with_suffix('.csv'))
    assert destination.read_bytes() == b'previous workbook'
    assert 'https://example.test' in Path(result).read_text(encoding='utf-8')


@pytest.mark.parametrize('count', [0, 1, 20, 21, 45])
def test_complete_findings_data_survives_pagination_and_escapes_untrusted_text(tmp_path, count):
    payload = '</script><img src=x onerror=alert(1)>&'
    issues = [{'severity': 'warning', 'category': 'SEO', 'message': f'Finding {i}: {payload}',
               'count': 1, 'pages_affected': 1, 'pages': [f'https://example.test/{i}'],
               'evidence': {'all_items': list(range(25))}, 'url': f'https://example.test/{i}'}
              for i in range(count)]
    data = {'pages': [], 'issues': {'aggregated': issues}}
    exporter = ExportManager({'output_directory': str(tmp_path)})
    output = Path(exporter.export(data, 'html')).read_text(encoding='utf-8')
    soup = BeautifulSoup(output, 'html.parser')
    assert json.loads(soup.select_one('#findings-data').string) == issues
    assert soup.select_one('#findings-search')['type'] == 'search'
    assert soup.select_one('#findings-severity')
    assert soup.select_one('#findings-next')
    assert not soup.select('[onerror]')
    assert not soup.find('img')


@pytest.fixture
def page():
    html = '''<html lang="en"><head><title>A useful title for the local export fixture</title>
    <meta name="description" content="A complete description preserved by every export format.">
    <meta name="viewport" content="width=device-width, initial-scale=1"></head>
    <body><h1>Export fixture</h1><p>A short local fixture for analysis.</p>
    <img src="/fixture-image.png"></body></html>'''
    return asyncio.run(SEOAnalyzer().analyze_page({
        'url': 'https://example.test/fixture', 'status_code': 200, 'load_time': 0.125,
        'content_length': len(html.encode()), 'headers': {'Content-Type': 'text/html'},
        'soup': BeautifulSoup(html, 'html.parser'),
    }))


@pytest.fixture
def exporter(tmp_path):
    return ExportManager({'output_directory': str(tmp_path)})


@pytest.mark.parametrize('template', ['report', 'enhanced', 'optimized'])
@pytest.mark.parametrize('kind', ['single', 'site'])
def test_actual_results_render_all_templates(tmp_path, page, template, kind):
    data = page if kind == 'single' else SEOAnalyzer().generate_site_report([page])
    before = copy.deepcopy(data)
    exporter = ExportManager({'output_directory': str(tmp_path), 'html_template': template})
    path = Path(exporter.export(data, 'html'))
    soup = BeautifulSoup(path.read_text(encoding='utf-8'), 'html.parser')
    text = soup.get_text(' ', strip=True)
    assert 'SEO Analysis' in text or 'TFQ0SEO Analysis' in text
    assert 'Not measured' in text or str(round(page['overall_score'])) in text
    assert "'recommendation':" not in text
    assert data == before
    prepared = exporter._prepare_html_data(data)
    assert prepared['category_scores']['seo'] == page['seo']['score']
    if kind == 'site':
        assert prepared['aggregated_issues'] == data['issues']['aggregated']
        assert prepared['aggregation_stats'] == data['issues']['stats']


@pytest.mark.parametrize('kind', ['single', 'site'])
def test_actual_results_export_csv_excel_json(tmp_path, exporter, page, kind):
    data = page if kind == 'single' else SEOAnalyzer().generate_site_report([page])
    original = copy.deepcopy(data)
    csv_path = Path(exporter.export(data, 'csv'))
    with csv_path.open(encoding='utf-8', newline='') as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 1
    assert rows[0]['title'] == page['seo']['data']['title']['text']
    assert rows[0]['description'] == page['seo']['data']['description']['text']
    assert float(rows[0]['seo_score']) == page['seo']['score']
    workbook = load_workbook(exporter.export(data, 'xlsx'))
    assert workbook['Summary']['B3'].value == pytest.approx(page['overall_score'])
    assert workbook['Pages'].max_row == 2
    assert workbook['Issues'].max_row > 1
    assert workbook['Recommendations'].max_row > 1
    workbook.close()
    json_path = Path(exporter.export(data, 'json'))
    assert json.loads(json_path.read_text(encoding='utf-8')) == json.loads(json.dumps(data))
    assert data == original


@pytest.mark.parametrize('template', ['report', 'enhanced', 'optimized', 'fallback'])
def test_untrusted_page_content_is_inert(tmp_path, page, template):
    payload = '<img src=x onerror=alert(1)>'
    html = f'<html lang="{payload}"><head><title>Fixture</title></head><body>Fixture</body></html>'
    data = asyncio.run(SEOAnalyzer().analyze_page({
        'url': 'https://example.test/hostile', 'status_code': 200, 'load_time': 0.1,
        'headers': {}, 'soup': BeautifulSoup(html, 'html.parser'),
    }))
    observed = next(issue for issue in data['issues']
                    if payload in json.dumps(issue.get('evidence', {}), ensure_ascii=False))
    # Keep the analyzer's evidence and also exercise imported message escaping,
    # including the minimal fallback template that does not display rule metadata.
    observed['message'] = 'Imported report message: ' + payload
    # The imported report path also accepts untrusted URL strings.
    data['url'] = 'javascript:alert(1)'
    exporter = ExportManager({'output_directory': str(tmp_path), 'html_template': 'report' if template == 'fallback' else template})
    if template == 'fallback':
        exporter.jinja_env.loader = FileSystemLoader(str(tmp_path / 'missing-templates'))
    path = Path(exporter.export(data, 'html'))
    rendered = path.read_text(encoding='utf-8')
    soup = BeautifulSoup(rendered, 'html.parser')
    assert not soup.select('[onerror]')
    assert not soup.find('img')
    assert all(not str(link.get('href', '')).lower().startswith('javascript:') for link in soup.find_all('a'))
    assert payload in soup.get_text()
    assert 'innerHTML' not in rendered


@pytest.mark.parametrize('value', ['=1+1', '+SUM(A1:A2)', '-1+2', '@SUM(A1)', '\t=1+1', '\r=1+1'])
def test_spreadsheet_formula_injection_is_neutralized(exporter, page, value):
    page['seo']['data']['title']['text'] = value
    page['issues'][0]['message'] = value
    page['recommendations'] = [{'recommendation': value, 'priority': 'high', 'impact': {'value': value}}]
    with Path(exporter.export(page, 'csv')).open(encoding='utf-8', newline='') as stream:
        row = next(csv.DictReader(stream))
    assert row['title'] == "'" + value
    workbook = load_workbook(exporter.export(page, 'xlsx'), data_only=False)
    for sheet in workbook:
        for cells in sheet:
            assert all(cell.data_type != 'f' for cell in cells)
    # XML backends may preserve or normalize carriage returns in inline strings.
    def normalize_line_endings(cell_value):
        if isinstance(cell_value, str):
            return cell_value.replace('\r\n', '\n').replace('\r', '\n')
        return cell_value

    excel_value = normalize_line_endings("'" + value)
    assert normalize_line_endings(workbook['Recommendations']['A2'].value) == excel_value
    assert any(normalize_line_endings(c.value) == excel_value for row in workbook['Issues'] for c in row)
    workbook.close()


@pytest.mark.parametrize('template', ['report', 'enhanced', 'optimized'])
def test_failed_and_unmeasured_results_do_not_become_zero(tmp_path, template):
    exporter = ExportManager({'output_directory': str(tmp_path), 'html_template': template})
    data = {'url': 'https://example.test/failed', 'error': 'Connection failed', 'overall_score': None,
            'performance': {'score': None, 'status': 'not_measured', 'data': {'metrics': {'lcp': None, 'fid': None}}},
            'seo': {'score': None, 'error': 'SEO analysis failed'}}
    path = Path(exporter.export(data, 'html'))
    soup = BeautifulSoup(path.read_text(encoding='utf-8'), 'html.parser')
    text = soup.get_text(' ', strip=True)
    assert 'Not measured' in text
    assert 'Connection failed' in text
    assert 'SEO analysis failed' in text
    assert '0/100' not in text
    assert all(value is None for value in exporter._prepare_html_data(data)['category_scores'].values())


@pytest.mark.parametrize('template', ['report', 'enhanced', 'optimized'])
def test_every_page_survives_large_report_export(tmp_path, page, template):
    pages = [dict(copy.deepcopy(page), url=f'https://example.test/page-{index}') for index in range(505)]
    failed = {'url': 'https://example.test/failed', 'error': 'Timed out', 'overall_score': None}
    report = SEOAnalyzer().generate_site_report(pages + [failed])
    exporter = ExportManager({'output_directory': str(tmp_path), 'html_template': template})
    prepared = exporter._prepare_html_data(report)
    assert len(prepared['pages_summary']) == 506
    assert prepared['pages_summary'][-1]['error'] == 'Timed out'
    html = Path(exporter.export(report, 'html')).read_text(encoding='utf-8')
    if template == 'optimized':
        embedded = re.search(r'const pagesData = (.*);', html).group(1)
        assert len(json.loads(embedded)) == 506
    else:
        soup = BeautifulSoup(html, 'html.parser')
        assert len(soup.select('.pages-table tbody tr')) == 506
    with Path(exporter.export(report, 'csv')).open(encoding='utf-8', newline='') as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 506
    assert rows[-1]['error'] == 'Timed out'
    workbook = load_workbook(exporter.export(report, 'xlsx'), read_only=True)
    assert workbook['Pages'].max_row == 507
    workbook.close()
    archived = json.loads(Path(exporter.export(report, 'json')).read_text(encoding='utf-8'))
    assert len(archived['pages']['detailed']) == 505
    assert len(archived['pages']['failed']) == 1


def test_export_preserves_missing_measurements_and_full_text(exporter, page):
    title = 'A long title ' * 30
    page['seo']['data']['title']['text'] = title
    page['content']['data']['readability'] = {'flesch_reading_ease': None}
    page['performance']['score'] = None
    with Path(exporter.export(page, 'csv')).open(encoding='utf-8', newline='') as stream:
        row = next(csv.DictReader(stream))
    assert row['title'] == title
    assert row['flesch_reading_ease'] == ''
    assert row['performance_score'] == ''


def test_fallback_to_csv_when_excel_extra_is_missing(exporter, page, monkeypatch):
    monkeypatch.setattr('tfq0seo.exporters.base.OPENPYXL_AVAILABLE', False)
    path = Path(exporter.export(page, 'xlsx'))
    assert path.suffix == '.csv'
    assert path.exists()


def test_enhanced_separates_observed_zero_from_unmeasured_metrics(tmp_path, page):
    exporter = ExportManager({'output_directory': str(tmp_path), 'html_template': 'enhanced'})
    page['overall_score'] = 0
    page['performance'] = {'score': None, 'data': {
        'measurement_status': 'browser_not_measured',
        'metrics': {'html_fetch_time': 0, 'content_size_mb': 0, 'lcp': None, 'inp': None},
    }}
    path = Path(exporter.export(page, 'html'))
    soup = BeautifulSoup(path.read_text(encoding='utf-8'), 'html.parser')
    metrics = {card.select_one('.metric-label').get_text(strip=True):
               card.select_one('.metric-value').get_text(strip=True)
               for card in soup.select('#performance .metric-card')}
    assert metrics['HTML fetch time'] == '0.00s'
    assert metrics['HTML size'] == '0.00MB'
    assert metrics['Largest Contentful Paint'] == 'Not measured'
    assert metrics['Interaction to Next Paint'] == 'Not measured'
    assert '0/100' in soup.get_text()


@pytest.mark.parametrize('url', ['data:text/html,<script>alert(1)</script>', '//example.test/path', 'https:\n//example.test'])
def test_report_rejects_unsafe_or_ambiguous_link_schemes(tmp_path, url):
    exporter = ExportManager({'output_directory': str(tmp_path), 'html_template': 'report'})
    path = Path(exporter.export({'url': url}, 'html'))
    soup = BeautifulSoup(path.read_text(encoding='utf-8'), 'html.parser')
    assert not soup.select('.url-info a')


@pytest.mark.parametrize(('pattern', 'expected_name'), [
    ('audit-{domain}.{format}', 'audit-example.test.csv'),
    ('{format}-for-{domain}', 'csv-for-example.test.csv'),
])
def test_configured_filename_patterns(tmp_path, page, pattern, expected_name):
    exporter = ExportManager({'output_directory': str(tmp_path), 'filename_pattern': pattern})
    output = Path(exporter.export(page, 'csv'))
    assert output == tmp_path / expected_name
    assert output.is_file()


def test_filename_timestamp_is_available_for_nested_site_reports(tmp_path, page):
    exporter = ExportManager({'output_directory': str(tmp_path), 'filename_pattern': '{domain}-{timestamp}.{format}'})
    report = SEOAnalyzer().generate_site_report([page])
    output = Path(exporter.export(report, 'json'))
    assert re.fullmatch(r'example\.test-\d{8}_\d{6}_\d{6}\.json', output.name)
    assert output.parent == tmp_path


@pytest.mark.parametrize('pattern', [
    '{unknown}', '{domain.__class__}', '{domain[0]}', '{domain!r}', '{domain:>30}',
    '{domain', '../{domain}', '..\\{domain}', '/outside/{domain}', 'C:\\outside\\{domain}', '',
    'audit:stream', 'NUL', 'CON', 'report\nname',
])
def test_invalid_filename_patterns_are_rejected_before_writing(tmp_path, page, monkeypatch, pattern):
    exporter = ExportManager({'output_directory': str(tmp_path), 'filename_pattern': pattern})

    def unexpected_write(*args, **kwargs):
        raise AssertionError('An invalid filename reached the writer')

    # Intercept dispatch so a validation regression cannot write devices or NTFS streams.
    monkeypatch.setattr(exporter, 'export_json', unexpected_write)
    with pytest.raises(ValueError):
        exporter.export(page, 'json')
    assert list(tmp_path.iterdir()) == []


def test_explicit_output_path_takes_precedence_over_filename_pattern(tmp_path, page):
    exporter = ExportManager({'output_directory': str(tmp_path), 'filename_pattern': '{unknown}'})
    output = tmp_path / 'requested.json'
    assert exporter.export(page, 'json', str(output)) == str(output)
    assert output.is_file()


@pytest.mark.parametrize('value', [float('nan'), float('inf'), float('-inf')])
def test_nonfinite_json_values_are_rejected_without_truncating_existing_report(tmp_path, exporter, value):
    output = tmp_path / 'preserved.json'
    previous = '{"previous_report": true}\n'
    output.write_text(previous, encoding='utf-8')
    data = {'url': 'https://example.test/', 'measurements': [{'value': value}]}
    with pytest.raises(ValueError):
        exporter.export(data, 'json', str(output))
    assert output.read_text(encoding='utf-8') == previous
