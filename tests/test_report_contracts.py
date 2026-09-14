"""Shared producer/import contracts and export validation boundaries."""

import asyncio
import copy
import csv
import json
from pathlib import Path

import pytest
from bs4 import BeautifulSoup
from click.testing import CliRunner
from openpyxl import load_workbook

from tfq0seo.cli import cli
from tfq0seo.core.app import SEOAnalyzer
from tfq0seo.core.models import ContractError
from tfq0seo.core.report_contracts import validate_page_result, validate_report, validate_site_report
from tfq0seo.exporters.base import ExportManager


@pytest.fixture
def page():
    html = '<html lang="en"><title>Example title</title><h1>Example heading</h1><p>Useful words for this local fixture.</p></html>'
    return asyncio.run(SEOAnalyzer().analyze_page({'url': 'https://example.test/page', 'status_code': 200,
        'soup': BeautifulSoup(html, 'html.parser'), 'content_length': len(html), 'headers': {}}))


def test_real_generated_envelopes_validate_strictly_without_mutation(page):
    report = SEOAnalyzer().generate_site_report([page])
    for value, validator in ((page, validate_page_result), (report, validate_site_report)):
        original = copy.deepcopy(value)
        assert validator(value, strict=True) is value
        assert validate_report(value, strict=True) is value
        assert value == original


def test_minimal_error_site_and_failed_page_remain_exportable():
    site = {'schema_version': '1.0', 'status': 'error', 'error': 'No pages to analyze'}
    assert validate_site_report(site, strict=True) is site
    assert validate_report(site, strict=True) is site
    page = asyncio.run(SEOAnalyzer().analyze_page({'url': 'https://example.test/failure', 'status_code': 404, 'error': 'Not found'}))
    assert validate_page_result(page, strict=True) is page


def test_imported_unknown_rules_and_untrusted_strings_are_not_rewritten():
    archived = {'schema_version': 'historical', 'url': 'javascript:alert(1)',
        'issues': [{'rule_id': 'removed.old_rule', 'rule_version': 'old', 'message': '<img src=x onerror=alert(1)>',
                    'evidence': {'html': '<script>example</script>'}, 'fix': 'Original archived guidance'}],
        'recommendations': ['Original recommendation'], 'page_facts': {'future_key': ['additive', None]}}
    original = copy.deepcopy(archived)
    assert validate_report(archived) is archived
    assert archived == original


@pytest.mark.parametrize('pages', [
    [{'url': 'https://example.test/', 'score': None, 'issues': {'total': 3}}],
    {'summary': [{'url': 'https://example.test/', 'score': 75, 'issues': {'warning': 2}}]},
    {'detailed': [{'url': 'https://example.test/', 'seo': {'score': 80}}], 'failed': [], 'skipped': []},
])
def test_historical_page_shapes_and_missing_optional_fields_are_supported(pages):
    report = {'pages': pages, 'recommendations': ['Legacy guidance']}
    assert validate_report(report) is report


def test_strict_site_can_contain_compatible_legacy_pages(page):
    report = SEOAnalyzer().generate_site_report([page])
    report['pages']['detailed'] = [{'url': page['url'], 'overall_score': 70}]
    assert validate_site_report(report, strict=True) is report


def test_site_generation_accepts_legacy_pages_without_http_measurements():
    report = SEOAnalyzer().generate_site_report([{'url': 'https://example.test/legacy', 'overall_score': 70, 'issues': []}])
    assert report['pages']['summary'][0]['status_code'] is None
    assert validate_site_report(report, strict=True) is report


@pytest.mark.parametrize('payload,path', [
    ({'seo': None}, '$.seo'),
    ({'seo': {'data': None}}, '$.seo.data'),
    ({'content': {'data': {'metrics': []}}}, '$.content.data.metrics'),
    ({'content': {'data': {'metrics': {'word_count': 'many'}}}}, '$.content.data.metrics.word_count'),
    ({'content': {'data': {'readability': None}}}, '$.content.data.readability'),
    ({'seo': {'data': {'title': {'text': []}}}}, '$.seo.data.title.text'),
    ({'seo': {'data': {'structured_data': None}}}, '$.seo.data.structured_data'),
    ({'issues': [None]}, '$.issues[0]'),
    ({'issues': [{'severity': []}]}, '$.issues[0].severity'),
    ({'issue_counts': {'total': True}}, '$.issue_counts.total'),
    ({'overall_score': '90'}, '$.overall_score'),
    ({'overall_score': 101}, '$.overall_score'),
    ({'scores': None}, '$.scores'),
    ({'scores': {'categories': {'seo': []}}}, '$.scores.categories.seo'),
    ({'recommendations': [False]}, '$.recommendations[0]'),
    ({'recommendations': {'specific': [7]}}, '$.recommendations.specific[0]'),
    ({'pages': [{'url': 'https://example.test/'}, 7]}, '$.pages[1]'),
    ({'pages': {'failed': {}}}, '$.pages.failed'),
    ({'pages': {'detailed': [{'url': 7}]}}, '$.pages.detailed[0].url'),
    ({'pages': [], 'issues': {'aggregated': [None]}}, '$.issues.aggregated[0]'),
])
def test_present_malformed_structure_is_rejected_with_its_path(payload, path):
    with pytest.raises(ContractError) as caught:
        validate_report(payload)
    assert path in str(caught.value)


@pytest.mark.parametrize('value', [float('nan'), float('inf'), float('-inf'), object(), {'unsupported'}, (1, 2)])
def test_nonportable_nested_extension_values_are_rejected(value):
    with pytest.raises(ContractError) as caught:
        validate_report({'url': 'https://example.test/', 'future_extension': {'value': value}})
    assert 'future_extension' in str(caught.value)


def test_strict_envelope_missing_fields_report_the_producer_path():
    with pytest.raises(ContractError) as caught:
        validate_page_result({'url': 'https://example.test/'}, strict=True, path='$.pages[0]')
    assert '$.pages[0].schema_version' in str(caught.value)
    with pytest.raises(ContractError) as caught:
        validate_site_report({'schema_version': '1.0', 'status': 'complete'}, strict=True)
    assert '$.summary' in str(caught.value)


def test_strict_failed_page_cannot_claim_a_measured_score(page):
    page.update(status='error', error='Fetch failed', overall_score=100)
    with pytest.raises(ContractError) as caught:
        validate_page_result(page, strict=True)
    assert '$.overall_score' in str(caught.value)
    page['overall_score'] = None
    assert validate_page_result(page, strict=True) is page


def test_strict_failed_page_requires_an_explanation(page):
    page.update(status='error', overall_score=None)
    with pytest.raises(ContractError) as caught:
        validate_page_result(page, strict=True)
    assert '$.error' in str(caught.value)


@pytest.mark.parametrize('format', ['json', 'csv', 'xlsx', 'html'])
@pytest.mark.parametrize('direct', [False, True])
def test_all_writer_entry_points_reject_invalid_payload_before_truncation(tmp_path, format, direct):
    exporter = ExportManager({'output_directory': str(tmp_path)})
    destination = tmp_path / ('existing.' + format)
    destination.write_bytes(b'previous report')
    payload = {'url': 'https://example.test/', 'content': {'data': {'metrics': None}}}
    with pytest.raises(ContractError):
        if direct:
            getattr(exporter, 'export_' + format)(payload, destination)
        else:
            exporter.export(payload, format, str(destination))
    assert destination.read_bytes() == b'previous report'


@pytest.mark.parametrize('literal', ['NaN', 'Infinity', '-Infinity', '1e400'])
def test_cli_rejects_nonfinite_json_before_creating_output(tmp_path, monkeypatch, literal):
    monkeypatch.chdir(tmp_path)
    source = tmp_path / 'source.json'
    destination = tmp_path / 'report.csv'
    source.write_text('{"url":"https://example.test/","overall_score":' + literal + '}', encoding='utf-8')
    result = CliRunner().invoke(cli, ['export', '-i', str(source), '-f', 'csv', '-o', str(destination)])
    assert result.exit_code != 0
    assert 'finite' in result.output.lower()
    assert not destination.exists()


@pytest.mark.parametrize('site', [False, True])
def test_real_word_count_reaches_csv_and_excel(tmp_path, page, site):
    expected = page['content']['data']['metrics']['word_count']
    report = SEOAnalyzer().generate_site_report([page]) if site else page
    exporter = ExportManager({'output_directory': str(tmp_path)})
    with Path(exporter.export(report, 'csv')).open(encoding='utf-8', newline='') as stream:
        assert int(next(csv.DictReader(stream))['word_count']) == expected
    workbook = load_workbook(exporter.export(report, 'xlsx'), read_only=True)
    try:
        sheet = workbook['Pages']
        headers = [cell.value for cell in next(sheet.iter_rows())]
        values = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
        assert values[headers.index('Word Count')] == expected
    finally:
        workbook.close()


def test_real_word_count_reaches_enhanced_html(tmp_path, page):
    exporter = ExportManager({'output_directory': str(tmp_path), 'html_template': 'enhanced'})
    output = Path(exporter.export(page, 'html'))
    soup = BeautifulSoup(output.read_text(encoding='utf-8'), 'html.parser')
    values = {card.select_one('.metric-label').get_text(strip=True): card.select_one('.metric-value').get_text(strip=True)
              for card in soup.select('#content .metric-card')}
    assert values['Word Count'] == str(page['content']['data']['metrics']['word_count'])


def test_legacy_flat_word_count_and_explicit_unknown_nested_count_are_preserved(tmp_path):
    exporter = ExportManager({'output_directory': str(tmp_path)})
    legacy = {'url': 'https://example.test/', 'content': {'data': {'word_count': 12}}}
    assert exporter._flatten_page_data(legacy)['word_count'] == 12
    legacy['content']['data']['metrics'] = {'word_count': None}
    assert exporter._flatten_page_data(legacy)['word_count'] is None
    assert exporter._prepare_html_data(legacy)['content_data']['word_count'] is None


def test_partial_legacy_summary_and_string_recommendations_still_render(tmp_path):
    report = {'pages': [{'url': 'https://example.test/', 'score': 75}],
              'recommendations': {'executive': {'overview': {'overall_health': 'Archived assessment'}},
                                  'specific': ['Preserve this legacy recommendation']},
              'aggregation_stats': {'unique_issues': 1}}
    original = copy.deepcopy(report)
    output = ExportManager({'output_directory': str(tmp_path)}).export(report, 'html')
    text = BeautifulSoup(Path(output).read_text(encoding='utf-8'), 'html.parser').get_text(' ', strip=True)
    assert 'Archived assessment' in text
    assert 'Preserve this legacy recommendation' in text
    assert report == original
